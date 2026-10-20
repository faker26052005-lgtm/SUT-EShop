# -*- coding: utf-8 -*-
"""
fill_ai_audit_report.py
Fills [AI-02] - FIT@HCMUS - AI Audit Report_Vn.docx authentically based on the HW02 conversation.
"""

import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_PATH = r"C:\Coding\SoftwareTesting\Seminar\src\eshop-sut\hw02-domain-testing\[AI-02] - FIT@HCMUS - AI Audit Report_Vn.docx"

doc = docx.Document(DOCX_PATH)

# ==========================================
# 1. Table 0: Student Information
# ==========================================
t0 = doc.tables[0]
t0.rows[1].cells[1].text = "NGUYỄN HOÀNG ANH TÚ"
t0.rows[2].cells[1].text = "23120183"
t0.rows[3].cells[1].text = "23_3 / k23"
t0.rows[4].cells[1].text = "HW#02 (HW02-AI: Domain Testing & BVA)"
t0.rows[5].cells[1].text = "20/10/2026"
t0.rows[6].cells[1].text = "Google Antigravity Agent (Gemini 2.5 / Claude 3.5 Sonnet)"
t0.rows[7].cells[1].text = "[ x] Có  [ ] Không"

# Format text in Table 0
for row in t0.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

# ==========================================
# 2. Table 1: Audit Table
# ==========================================
t1 = doc.tables[1]

# Delete rows from index 1 down to the end (except row 0 header)
# We will then add the 6 authentic artifact rows.
while len(t1.rows) > 1:
    row_to_del = t1.rows[1]
    row_to_del._tr.getparent().remove(row_to_del._tr)

artifacts_data = [
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 10:20 20/10/2026\n"
            'Prompt: "Hiện tại trong .csv, Test data: Name=baseline có nghĩa là gì"'
        ),
        "output": (
            'AI giải thích "Name=baseline" là giá trị ngầm định trong kỹ thuật kiểm thử 1D domain testing, '
            'đại diện cho việc giữ nguyên các trường khác ở mức baseline hợp lệ.'
        ),
        "verdict": "INVALID",
        "reason": (
            "Slide S07 Test Case Engineering quy định: Test Data bắt buộc phải là các giá trị literal cụ thể (explicit literals), "
            'không được dùng placeholder giả định vì tester không thể nhập chuỗi "baseline" vào form web thực tế.'
        ),
        "revision": (
            'Yêu cầu AI loại bỏ toàn bộ chuỗi "baseline", thay thế bằng các giá trị thực tế cụ thể: '
            'Name="Nguyen Van A", Email="test@eshop.com", Password="Password123@"; '
            'cập nhật lại toàn bộ 28 test cases trong FR01_TestCases.csv.'
        )
    },
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 10:35 20/10/2026\n"
            'Prompt: "Tại sao có thể test được TC-FR01-D15 khi không có trường Confirm pw" -> "Sửa lại"'
        ),
        "output": (
            'Ban đầu AI sinh các bước thao tác yêu cầu nhập liệu vào ô "Confirm password" trên form Register để kiểm thử trùng khớp mật khẩu.'
        ),
        "verdict": "INVALID",
        "reason": (
            'Vi phạm nguyên lý ISTQB "Testing is context-dependent". AI tự suy diễn (hallucinate) sự tồn tại của ô Confirm Password, '
            'trong khi mã nguồn Register.jsx thực tế của EShop chỉ có 3 trường Name, Email, Password (BUG-FR01-02).'
        ),
        "revision": (
            'Chuyển đổi TC-FR01-D15 thành Test Case kiểm chứng lỗi thiết kế (Defect Verification), ghi nhận kết quả FAILED với quan sát: '
            '"Giao diện thiếu hoàn toàn ô Xác nhận mật khẩu theo đặc tả FR-01 dòng 22", bắt BUG-FR01-02.'
        )
    },
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 11:15 20/10/2026\n"
            'Prompt: "Áp dụng skill domain-boundary-testing tiếp tục cho FR09"'
        ),
        "output": (
            "AI phân tích miền giá trị mã giảm giá, sinh ma trận BVA cho các mã SAVE10 (10%), FIXED50K (50.000₫), FREESHIP (15.000₫) "
            "với Expected Result tính toán theo lý thuyết số học thông thường."
        ),
        "verdict": "INCOMPLETE",
        "reason": (
            "ISTQB FL §4.3 yêu cầu đối chiếu biên thực tế trên SUT. Khi chạy thực tế trên API, phát hiện SUT dính lỗi toán học chia/nhân sai "
            "(SAVE10 tính ngược làm tăng tiền) và cho phép giảm giá cố định vượt quá tổng đơn hàng (FIXED50K trên đơn 30.000₫)."
        ),
        "revision": (
            "Bổ sung AI Gap Analysis chỉ ra lỗi logic nghiệp vụ trong server.js, chuẩn hóa 31 Test Cases theo Slide S07 "
            "và lập 6 Bug Reports chi tiết (BUG-FR09-01 đến 06) theo chuẩn Slide S08."
        )
    },
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 11:45 20/10/2026\n"
            'Prompt: "Tiếp tục áp dụng skill domain-boundary-testing cho FR-15"'
        ),
        "output": (
            "AI sinh 30 test cases kiểm thử CRUD sản phẩm (Name, Price, Category) trên Web Admin nhưng ban đầu chỉ tập trung vào nghiệp vụ thông thường, "
            "bỏ qua kiểm thử an toàn chuỗi truy vấn."
        ),
        "verdict": "INCOMPLETE",
        "reason": (
            "Phân tích miền giá trị đối với trường văn bản (Text Domain) phải bao gồm các ký tự đặc biệt của cú pháp SQL "
            "(' OR '1'='1) để kiểm thử ranh giới xử lý injection của cơ sở dữ liệu."
        ),
        "revision": (
            "Bổ sung TC-FR15-D18 kiểm thử SQL Injection và TC-FR15-D19 kiểm thử JWT Access Control, phát hiện lỗ hổng Fatal SQL Injection "
            "tại server.js:150 (BUG-FR15-05) và lưu CSDL giá âm (BUG-FR15-03)."
        )
    },
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 12:20 20/10/2026\n"
            'Prompt: "Các test case phân hệ mobile đã được chạy thực sự chưa hay chỉ dựa trên phân tích mã nguồn" -> "Hướng 2"'
        ),
        "output": (
            "AI sinh báo cáo kiểm thử Mobile Cart nhưng ban đầu chưa làm rõ môi trường thực thi, ngầm định là test trên ứng dụng chạy thực tế như Web."
        ),
        "verdict": "INCOMPLETE",
        "reason": (
            "Môi trường kiểm thử trong Bug Report và Test Report bắt buộc phải chính xác và minh bạch. "
            "Thư mục frontend-mobile không có node_modules và không khởi chạy máy ảo Expo."
        ),
        "revision": (
            "Chọn \"Hướng 2\": Chèn khối tuyên ngôn minh bạch phương pháp luận kiểm thử hộp trắng (White-box Static Code Inspection trên App.js) "
            "kết hợp Backend API Integration vào đầu Feature_D_MobileCart_Report.md và Bug_Report.md; cập nhật trường Environment của BUG-MOB-01..06."
        )
    },
    {
        "prompt": (
            "Tool: Google Antigravity Agent\n"
            "Thời gian: 13:45 20/10/2026\n"
            'Prompt: "Có cách nào làm agent tự tạo thay vì thêm thủ công không" -> Cung cấp GitHub Token'
        ),
        "output": (
            "AI viết script auto_create_github_issues.js nhưng regex ban đầu (BUG-[A-Z0-9]+) chỉ bắt được BUG-FR01 (thiếu dấu gạch nối thứ hai), "
            "khiến số bug phân tích ra bằng 0."
        ),
        "verdict": "INVALID",
        "reason": (
            "Lỗi cú pháp Regex parsing không khớp với quy ước đặt mã lỗi chuẩn của dự án (BUG-FR01-01)."
        ),
        "revision": (
            "Sửa regex thành ^(BUG-[A-Z0-9]+-[0-9]+), chạy script thành công tạo trọn vẹn 25 issues (#52 đến #76) trên repo GitHub "
            "kèm nhãn Labels và xuất bảng minh chứng vào Evidence/GitHub_Issues/README.md."
        )
    }
]

for art in artifacts_data:
    row = t1.add_row()
    row.cells[0].text = art["prompt"]
    row.cells[1].text = art["output"]
    row.cells[2].text = art["verdict"]
    row.cells[3].text = art["reason"]
    row.cells[4].text = art["revision"]
    for c_idx, cell in enumerate(row.cells):
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if c_idx == 2: # Verdict column centered
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    if art["verdict"] == "INVALID":
                        run.font.color.rgb = RGBColor(180, 0, 0)
                    elif art["verdict"] == "INCOMPLETE":
                        run.font.color.rgb = RGBColor(180, 100, 0)

# ==========================================
# 3. Table 2: Accuracy Summary
# ==========================================
t2 = doc.tables[2]
t2.rows[1].cells[1].text = "6"
t2.rows[1].cells[2].text = "100%"

t2.rows[2].cells[1].text = "0"
t2.rows[2].cells[2].text = "0%"

t2.rows[3].cells[1].text = "3"
t2.rows[3].cells[2].text = "50%"

t2.rows[4].cells[1].text = "3"
t2.rows[4].cells[2].text = "50%"

for row in t2.rows:
    for c_idx, cell in enumerate(row.cells):
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if c_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ==========================================
# 4. Paragraphs: Section 5 & 6
# ==========================================
conclusion_text = (
    "Qua quá trình thực hiện kiểm thử miền (Domain Testing) và phân tích giá trị biên (BVA) trên hệ thống EShop, "
    "tôi nhận thấy AI cực kỳ mạnh mẽ trong việc phân rã ma trận kiểm thử, tự động hóa tính toán các điểm biên (Min-, Min, Max, Max+), "
    "sinh mã kịch bản tự động hóa API và định dạng báo cáo chuẩn công nghiệp rất nhanh chóng. "
    "Tuy nhiên, AI thường có xu hướng \"happy-path bias\", tự suy đoán (hallucinate) các phần tử giao diện không tồn tại trên thực tế "
    "(như trường Confirm Password trên form đăng ký), dùng test data trừu tượng (baseline) và thiếu nhạy bén với các lỗi bảo mật hoặc mã lỗi "
    "cài cắm có chủ đích nếu không được chỉ dẫn kiểm thử hộp trắng.\n\n"
    "Khuyến nghị: CHỈ NÊN dùng AI như trợ thủ gia tốc thiết kế ca kiểm thử và tự động hóa tác vụ lặp lại. "
    "BẮT BUỘC người kiểm thử phải trực tiếp rà soát mã nguồn thực tế, kiểm chứng kết quả chạy thực nghiệm và minh bạch phương pháp luận thực thi."
)

disclosure_text = (
    '"Bộ test case và báo cáo kiểm thử này được sinh phiên bản đầu bởi Google Antigravity Agent (Gemini / Claude); '
    'tôi đã rà soát và chỉnh sửa loại bỏ toàn bộ dữ liệu giả định baseline, điều chỉnh kịch bản kiểm thử trường Confirm Password, '
    'bổ sung test case SQL Injection và minh bạch phương pháp luận hộp trắng cho phân hệ Mobile; '
    'toàn bộ 25 bugs thực tế trên mã nguồn EShop do tôi trực tiếp chỉ đạo AI cô lập và kiểm chứng. '
    'AI Audit Report chi tiết đính kèm ở Phụ lục A. Tôi cam đoan không dùng AI để sinh bất kỳ artifact nào thuộc danh mục bị cấm."'
)

# Replace placeholders in paragraphs
# Find section 5 and 6
for i, p in enumerate(doc.paragraphs):
    if "5. Kết luận — Khi nào nên / không nên dùng AI?" in p.text:
        # replace the underline lines (p+2 to p+7)
        if i + 2 < len(doc.paragraphs):
            doc.paragraphs[i + 2].text = conclusion_text
            # Clear lines i+3 to i+7
            for j in range(i + 3, min(i + 8, len(doc.paragraphs))):
                if "_____" in doc.paragraphs[j].text:
                    doc.paragraphs[j].text = ""

    if "6. Mandatory Disclosure (dán nguyên văn)" in p.text:
        if i + 1 < len(doc.paragraphs):
            doc.paragraphs[i + 1].text = disclosure_text
            doc.paragraphs[i + 1].runs[0].font.italic = True

# ==========================================
# 5. Table 3: Signatures
# ==========================================
t3 = doc.tables[3]
t3.rows[0].cells[1].text = "Nguyễn Hoàng Anh Tú"
t3.rows[1].cells[1].text = "23120183"
t3.rows[2].cells[1].text = "23_3 / k23"
t3.rows[3].cells[1].text = "CS423 / CSC13003 – Kiểm chứng Phần mềm"
t3.rows[4].cells[1].text = "Dr. Lâm Quang Vũ / TS. Trần Duy Hoàng / ThS. Trần Thị Bích Hạnh"
t3.rows[5].cells[1].text = "20/10/2026"
t3.rows[6].cells[1].text = "Nguyễn Hoàng Anh Tú"

for row in t3.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

doc.save(DOCX_PATH)
print("SUCCESS: Đã điền xong file DOCX:", DOCX_PATH)
